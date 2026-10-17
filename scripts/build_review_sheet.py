"""Build the expert review workbook for the v2 catalogue labels (event-focused criteria).

Input : annotation/catalogue_labels_v2_proposed.csv  (sent_id, sentence, type, label, note)
Output: annotation/catalogue_review_v2.xlsx           (expert edits final_type / final_label / reviewer_notes)
An existing review workbook is archived first, never overwritten.
"""
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from annotation_sheets import F, _cell, _header, _notes  # noqa: E402
from lasec.utils import ROOT, load_config, path  # noqa: E402
from prelabel_catalogue import archive  # noqa: E402

TYPES = ["event", "compound", "instruction", "mitigation", "possible", "info"]
TYPE_ORDER = {t: i for i, t in enumerate(TYPES)}


def main():
    cfg = load_config()
    lab = pd.read_csv(ROOT / "annotation" / "catalogue_labels_v2_proposed.csv", keep_default_na=False)
    occ = pd.read_csv(path(cfg, "processed") / "catalogue_occurrences.csv", keep_default_na=False)
    n_in_msg = occ.groupby("message").sentence.transform("size")
    occ["ctx"] = occ.message.where(n_in_msg > 1, "")
    ctx = occ.groupby("sentence").agg(source_message=("ctx", "first"),
                                      problem_type=("problem_type", lambda x: "/".join(sorted(set(x)))))
    lab = lab.join(ctx, on="sentence")
    lab["group_size"] = lab.groupby("label").sentence.transform("size")
    lab["needs_check"] = lab.note.str.contains("check").map({True: "Y", False: ""})
    lab["_t"] = lab["type"].map(TYPE_ORDER)
    lab = lab.sort_values(["_t", "label", "sentence"]).drop(columns="_t").reset_index(drop=True)

    out = ROOT / "annotation" / "catalogue_review_v2.xlsx"
    if out.exists():
        archive(out)
    wb = Workbook()
    ws = wb.active
    ws.title = "How to use"
    _notes(ws, [
        ("Catalogue labels v2 (event-focused)", "Proposed by Claude under the refined criteria (GUIDELINE.md v2). Please review."),
        ("Criteria", "1) Group EVENT sentences (state changes, actions, conditions) by their underlying event; instructions, "
                     "mitigations, possible events and information are not events.  2) A sentence with an event part and a "
                     "non-event part takes the event's label; a sentence with two or more events gets its own label (type "
                     "'compound').  3) Semantically similar sentences with the same underlying event share a label."),
        ("Where to edit", "Sheet 'Sentences', yellow columns: final_type, final_label (pre-filled with the proposal), reviewer_notes."),
        ("Order", "Rows are sorted by type (event, compound, instruction, mitigation, possible, info), then label, so each "
                  "group is contiguous."),
        ("Priority", "Filter needs_check = Y first: 30 decisions where the criteria allow two readings (the rationale says which)."),
        ("Non-event labels", "Non-events are still grouped by meaning (INSTR_ for instructions and mitigations, POSS_ for possible "
                             "events, INFO_ for information), so they can be included or excluded later without relabelling."),
        ("Expert decisions", "Your 7 edits of 12 Oct are kept (rationale 'expert decision' / 'expert ID')."),
        ("FPV IDs", "Catalogue sentences describing an FPV event reuse the FPV ID (sheet 'FPV events')."),
        ("Column 'changed'", "Formula: 'yes' when final_label or final_type differs from the proposal (correction rate)."),
        ("Example", "C0016 'Aircraft is in Attitude mode and hovering may be unstable.' -> event, AIRCRAFT_ATTITUDE_MODE "
                    "(possible consequence dropped)."),
    ])

    cols = ["sent_id", "sentence", "source_message", "problem_type", "type", "label", "group_size", "note",
            "needs_check", "final_type", "final_label", "changed", "reviewer_notes"]
    heads = ["sent_id", "sentence", "source message (context)", "problem_type (DroPTC)", "proposed type",
             "proposed label", "group size", "rationale", "needs_check", "final_type", "final_label", "changed",
             "reviewer_notes"]
    widths = [8, 50, 50, 14, 11, 34, 7, 44, 8, 12, 34, 8, 30]
    L = {c: get_column_letter(i) for i, c in enumerate(cols, 1)}
    ss = wb.create_sheet("Sentences")
    _header(ss, heads, widths)
    for r, row in enumerate(lab.itertuples(index=False), 2):
        d = row._asdict()
        for c, col in enumerate(cols, 1):
            if col == "final_type":
                _cell(ss, r, c, d["type"], edit=True)
            elif col == "final_label":
                _cell(ss, r, c, d["label"], edit=True)
            elif col == "reviewer_notes":
                _cell(ss, r, c, "", edit=True)
            elif col == "changed":
                _cell(ss, r, c, f'=IF(AND(TRIM({L["final_label"]}{r})={L["label"]}{r},'
                                 f'TRIM({L["final_type"]}{r})={L["type"]}{r}),"","yes")')
            else:
                _cell(ss, r, c, d[col])
    n = len(lab) + 1
    ss.freeze_panes = "C2"
    ss.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n}"
    dv = DataValidation(type="list", formula1='"' + ",".join(TYPES) + '"')
    ss.add_data_validation(dv)
    dv.add(f'{L["final_type"]}2:{L["final_type"]}{n}')
    nc, chg = L["needs_check"], L["changed"]
    ss.conditional_formatting.add(f"{nc}2:{nc}{n}", FormulaRule(formula=[f'${nc}2="Y"'], fill=PatternFill("solid", fgColor="F8CBAD")))
    ss.conditional_formatting.add(f"{chg}2:{chg}{n}", FormulaRule(formula=[f'${chg}2="yes"'], font=Font(name=F, bold=True, color="C00000")))

    ev = wb.create_sheet("Labels")
    _header(ev, ["proposed label", "type", "members", "sentences", "rows now labelled with this ID"], [34, 11, 8, 100, 14])
    g = lab.groupby("label", sort=False)
    fl = L["final_label"]
    for r, (e, sub) in enumerate(g, 2):
        for c, v in enumerate([e, sub["type"].iloc[0], len(sub), " | ".join(sub.sentence),
                               f"=COUNTIF(Sentences!${fl}$2:${fl}${n},A{r})"], 1):
            _cell(ev, r, c, v)
    ev.freeze_panes = "B2"
    ev.auto_filter.ref = f"A1:E{len(g) + 1}"

    fpv = pd.read_csv(path(cfg, "fpv_labels"))
    fe = wb.create_sheet("FPV events")
    _header(fe, ["FPV event ID (expert)", "proposed type (new)", "FPV sentences"], [30, 14, 100])
    for r, (e, sub) in enumerate(fpv.groupby(cfg["labels"]["fpv_column"]), 2):
        for c, v in enumerate([e, sub.sentence_type.iloc[0], " | ".join(sub.sentence)], 1):
            _cell(fe, r, c, v)
    for sh in wb.worksheets:
        sh.sheet_view.zoomScale = 90
    wb.save(out)
    print(f"wrote {out.relative_to(ROOT)}: {len(lab)} sentences, {lab.label.nunique()} labels, "
          f"{int((lab.needs_check == 'Y').sum())} to check")


if __name__ == "__main__":
    main()
