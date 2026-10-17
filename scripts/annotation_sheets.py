"""Excel helpers for the expert review workbook (scripts/build_review_sheet.py)."""
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = "Arial"
HDR = PatternFill("solid", fgColor="1F3864")
EDIT = PatternFill("solid", fgColor="FFF2CC")
BOX = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws, cols, widths, row=1):
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row, j, c)
        cell.fill, cell.font, cell.border = HDR, Font(name=F, bold=True, color="FFFFFF", size=10), BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w


def _cell(ws, r, c, v, edit=False, bold=False):
    cell = ws.cell(r, c, v)
    cell.font, cell.alignment, cell.border = Font(name=F, size=10, bold=bold), WRAP, BOX
    if edit:
        cell.fill = EDIT
    return cell


def _notes(ws, lines, width=120):
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = width
    for i, (a, b) in enumerate(lines, 1):
        ws.cell(i, 1, a).font = Font(name=F, size=11 if i == 1 else 10, bold=True, color="1F3864" if i == 1 else None)
        c = ws.cell(i, 2, b)
        c.font, c.alignment = Font(name=F, size=10), WRAP


def archive(file: Path) -> Path:
    """Move an existing sheet to annotation/archive/<name>_<timestamp>.xlsx (never overwrite work)."""
    dst = file.parent / "archive" / f"{file.stem}_{datetime.now():%Y%m%d-%H%M%S}{file.suffix}"
    dst.parent.mkdir(exist_ok=True)
    shutil.move(str(file), dst)
    return dst
